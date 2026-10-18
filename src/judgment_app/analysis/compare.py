"""比對檢核 Excel 與分析 JSON；以 --help 查看用法。"""

import argparse
from collections import Counter
import json
from pathlib import Path
import re
from tempfile import NamedTemporaryFile

import pandas as pd
from openpyxl import Workbook, load_workbook

from judgment_app.analysis.models import CrimeAnalysis, JudgmentAnalysis, LawReference
from judgment_app.case_number import CaseNumber
from judgment_app.paths import resolve_output_path

# 同名欄位以第一欄為準：Excel 後面的拘役等欄位屬法定刑上下限。
FIELD_MAP = {
    "裁判情形": "crime_jsst", "徒刑(YYMM)": "penalty", "拘役": "detention", "罰金": "fine",
    "累犯": "is_recidivist", "案由": "law_name", 
    "條": "article", "之": "article_sub", "條段": "article_part",
    "項": "paragraph", "項段": "paragraph_part", "款": "paragraph_sub",
    "未遂": "is_attempted", "幫助": "is_accessory",
}
IDENTITY_COLUMNS = {"裁判機關", "裁判年度", "裁判冠字", "裁判號數", "姓名"}
LAW_FIELDS = {"law_name", "article", "article_sub", "article_part", "paragraph", "paragraph_part", "paragraph_sub", "is_attempted", "is_accessory"}
PENALTIES = {"crime_jsst", "penalty", "detention", "fine"}



def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", default="data", help="Excel 檔案或資料夾（預設 data）")
    parser.add_argument("--json-dir", default="/Users/hsiaomt/Desktop/Judgments_101資料夾_最新格式", help="分析 JSON 檔案或資料夾，遞迴搜尋")
    parser.add_argument("--output", default="data/comparison_report.json")
    args = parser.parse_args()
    report = compare_files(resolve_output_path(args.excel), resolve_output_path(args.json_dir))
    destination = resolve_output_path(args.output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False))
    print(f"比對報告：{destination}")
    for output in report["excel_outputs"]:
        print(f"Excel 比對結果：{output['output']}")


def load_comparison_report(path: str | Path) -> dict:
    """讀取既有報告，驗證顯示所需欄位；不修改檔案。"""
    source = resolve_output_path(path)
    report = json.loads(source.read_text(encoding="utf-8-sig"))
    if not isinstance(report, dict) or not isinstance(report.get("rows"), list) or not isinstance(report.get("errors"), list):
        raise ValueError(f"不是有效的比對報告（需包含 rows、errors 清單）：{source}")
    for index, row in enumerate(report["rows"], 1):
        location = f"rows 第 {index} 筆"
        if not isinstance(row, dict) or not isinstance(row.get("status"), str) or not row["status"]:
            raise ValueError(f"{location} 缺少有效的 status")
        if row.get("reason") is not None and not isinstance(row["reason"], str):
            raise ValueError(f"{location} 的 reason 必須是字串")
        for field, required in (("differences", ("column",)), ("uncomparable", ("column", "reason"))):
            items = row.get(field, [])
            if not isinstance(items, list) or any(
                not isinstance(item, dict) or any(not isinstance(item.get(key), str) for key in required)
                for item in items
            ):
                raise ValueError(f"{location} 的 {field} 格式錯誤")
    for index, error in enumerate(report["errors"], 1):
        if not isinstance(error, dict) or not isinstance(error.get("error"), str):
            raise ValueError(f"errors 第 {index} 筆缺少錯誤訊息")
    # 摘要依實際內容計算，避免外部報告的過期統計影響顯示。
    report["summary"] = dict(Counter(row["status"] for row in report["rows"]))
    report["summary"].update(total=len(report["rows"]), errors=len(report["errors"]))
    return report


def compare_to_report(excel_path: Path, json_path: Path, destination: Path, *, progress=None) -> dict:
    """比對並回寫 Excel status，另存完整 JSON 報告。"""
    destination = resolve_output_path(destination).resolve()
    for source in (Path(excel_path), Path(json_path)):
        if source.is_file() and source.resolve() == destination:
            raise ValueError("比對報告不可覆寫來源檔案")
    if destination.suffix.lower() != ".json":
        raise ValueError("比對報告必須為 .json 檔案")
    if Path(json_path).is_dir() and destination.exists():
        data = json.loads(destination.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict) and "case_number" in data:
            raise ValueError("比對報告不可覆寫分析 JSON")
    report = compare_files(excel_path, json_path, progress=progress)
    if progress:
        progress(f"正在儲存比對報告：{destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def compare_files(excel_path: Path, json_path: Path, *, progress=None) -> dict:
    """回傳每列差異、未配對原因及未比較欄位，並將 status 寫入 Excel。"""
    report = {"rows": [], "errors": [], "sheets": [], "excel_outputs": []}
    if progress:
        progress("正在讀取分析 JSON…")
    index = _load_analyses(Path(json_path), report["errors"])
    files = _files(Path(excel_path), {".xls", ".xlsx"})
    if Path(excel_path).is_dir():
        files = [p for p in files if not re.search(r"\.status(?:-\d+)?\.xlsx$", p.name, re.IGNORECASE)]
    if not files:
        raise ValueError(f"找不到 Excel：{excel_path}")
    for file_index, path in enumerate(files, 1):
        if progress:
            progress(f"正在比對 Excel {file_index} / {len(files)}：{path.name}")
        try:
            statuses = {}
            legacy_tables = {} if path.suffix.lower() == ".xls" else None
            with pd.ExcelFile(path) as book:
                for sheet in book.sheet_names:
                    if progress:
                        progress(f"正在比對 {path.name}／{sheet}")
                    table = book.parse(sheet, header=None, dtype=object)
                    if legacy_tables is not None:
                        legacy_tables[sheet] = table
                    header = next((i for i, row in table.iterrows()
                                   if IDENTITY_COLUMNS <= {_text(x) for x in row}), None)
                    if header is None:
                        report["errors"].append({"file": str(path), "sheet": sheet, "error": "找不到裁判案號及姓名標題列"})
                        continue
                    columns = {}
                    for i, name in enumerate(table.iloc[header]):
                        columns.setdefault(_text(name), i)
                    # status 是前次比對的輸出，不參與欄位比對或空白列判斷。
                    columns.pop("status", None)
                    sheet_statuses = {}
                    statuses[sheet] = (header + 1, sheet_statuses)
                    report["sheets"].append({"file": str(path), "sheet": sheet,
                        "uncompared_columns": [c for c in columns if c and c not in FIELD_MAP and c not in IDENTITY_COLUMNS]})
                    for i in range(header + 1, len(table)):
                        row = {name: table.iat[i, col] for name, col in columns.items() if name}
                        if all(_blank(x) for x in row.values()):
                            continue
                        context = {"excel": str(path), "sheet": sheet, "row": i + 1, "name": _text(row.get("姓名"))}
                        try:
                            context.update(_compare_row(row, index))
                        except (ValueError, TypeError, KeyError) as exc:
                            context.update(status="unmatched", reason=str(exc))
                        report["rows"].append(context)
                        sheet_statuses[i + 1] = context["status"]
            if statuses:
                if progress:
                    progress(f"正在寫入 Excel 比對結果：{path.name}")
                output = _write_excel_statuses(path, statuses, legacy_tables=legacy_tables)
                report["excel_outputs"].append({"source": str(path), "output": str(output)})
                if progress:
                    progress(f"已儲存 Excel 比對結果：{output}")
        except Exception as exc:
            report["errors"].append({"file": str(path), "error": f"{type(exc).__name__}: {exc}"})
    report["summary"] = dict(Counter(row["status"] for row in report["rows"]))
    report["summary"].update(total=len(report["rows"]), errors=len(report["errors"]))
    return report


def _write_excel_statuses(path: Path, statuses: dict, *, legacy_tables=None) -> Path:
    """xlsx 原檔回寫；xls 以讀取的表格另存 xlsx，不覆寫既有結果檔。"""
    destination = path
    if legacy_tables is None:
        workbook = load_workbook(path)
    else:
        destination = path.with_name(f"{path.stem}.status.xlsx")
        number = 2
        while destination.exists():
            destination = path.with_name(f"{path.stem}.status-{number}.xlsx")
            number += 1
        workbook = Workbook()
        workbook.remove(workbook.active)
    temporary = None
    try:
        if legacy_tables is not None:
            for sheet, table in legacy_tables.items():
                worksheet = workbook.create_sheet(sheet)
                for row_number, row in enumerate(table.itertuples(index=False, name=None), 1):
                    for column, value in enumerate(row, 1):
                        if not pd.isna(value):
                            cell = worksheet.cell(row_number, column, value)
                            if isinstance(value, str):
                                cell.data_type = "s"
        for sheet, (header, row_statuses) in statuses.items():
            worksheet = workbook[sheet]
            column = next((cell.column for cell in worksheet[header]
                           if _text(cell.value) == "status"), worksheet.max_column + 1)
            worksheet.cell(header, column, "status")
            for row_number in range(header + 1, worksheet.max_row + 1):
                worksheet.cell(row_number, column).value = row_statuses.get(row_number)
        with NamedTemporaryFile(dir=path.parent, prefix="~$status-", suffix=".xlsx", delete=False) as file:
            temporary = Path(file.name)
        workbook.save(temporary)
        temporary.replace(destination)
        return destination
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _compare_row(row: dict, index: dict[CaseNumber, list[tuple[Path, JudgmentAnalysis]]]) -> dict:
    try:
        case = CaseNumber(_text(row["裁判機關"]), _integer(row["裁判年度"]),
                          _text(row["裁判冠字"]), _integer(row["裁判號數"]))
    except (ValueError, TypeError, KeyError) as exc:
        return {"status": "unmatched", "reason": str(exc)}
    base = {"case_number": case.to_string()}
    documents = index.get(case, [])
    if not documents:
        return {**base, "status": "unmatched", "reason": "判決不存在"}
    results = []
    for source, data in documents:
        context = {**base, "json": str(source)}
        try:
            defendants = [d for d in data.defendants if _matches_defendant_name(row["姓名"], d.name)]
            if len(defendants) != 1:
                results.append({**context, "status": "unmatched",
                        "reason": "被告姓名無法唯一配對"})
                continue
            defendant = defendants[0]

            # 只比對有罪的 group，若無則視為無法配對。
            guilty_groups = [g for g in defendant.execution_groups if g.executed_jsst in {"科刑", "免刑"}]
            if not guilty_groups:
                results.append({**context, "status": "unmatched", "reason": "找不到有罪的 group"})
                continue
            guilty_crime_ids = {crime_id for g in guilty_groups for crime_id in g.crime_ids}

            # 每列只轉換一次；所有候選均依法條及各罪刑度定位。
            values = {field: _excel_col_to_val(column, row[column])
                      for column, field in FIELD_MAP.items() if column in row}
            excel_law = LawReference(**{
                **{field: values[field] for field in LAW_FIELDS if field in values},
                "is_attempted": values.get("is_attempted", False),
                "is_accessory": values.get("is_accessory", False),
            })
            candidates = [(crime, law) for crime in defendant.crimes
                          if crime.crime_id in guilty_crime_ids for law in (crime.laws or [])
                          if _matches_candidate(values, crime, excel_law, law)]
            if not candidates:
                results.append({**context, "status": "unmatched",
                        "reason": "法條／刑度無符合候選", "candidate_crime_ids": []})
                continue
            # 多個符合候選也視為定位成功，保留全部候選 ID；累犯由被告資料比較。
            context["candidate_crime_ids"] = list(dict.fromkeys(c.crime_id for c, _ in candidates))
            if "is_recidivist" in values and values["is_recidivist"] != defendant.is_recidivist:
                results.append({**context, "status": "unmatched", "reason": "累犯不合"})
            else:
                results.append({**context, "status": "equal"})
        except (ValueError, TypeError, KeyError) as exc:
            results.append({**context, "status": "unmatched", "reason": str(exc)})
    if len(results) == 1:
        return results[0]
    matched = [result for result in results if result["status"] == "equal"]
    if not matched:
        return {**base, "status": "unmatched",
                "reason": "所有判決均無符合候選，詳見各判決結果", "candidate_crime_ids": [],
                "document_results": results}

    # 同案號只要有一份完全符合，就視為配對成功；保留每份判決的結果供查閱。
    return {**base, "status": "equal",
            "candidate_crime_ids": list(dict.fromkeys(cid for result in matched for cid in result["candidate_crime_ids"])),
            "matched_candidates": [{"json": result["json"], "crime_id": cid}
                                   for result in matched for cid in result["candidate_crime_ids"]],
            "document_results": results}


def _matches_candidate(values: dict, crime: CrimeAnalysis, excel_law: LawReference,
                       json_law: LawReference) -> bool:
    """以轉換後的 Excel 欄位定位，刑度固定使用 CrimeAnalysis 的各罪宣告刑。"""
    if excel_law != json_law:
        return False
    for field, expected in values.items():
        if field not in PENALTIES:
            continue
        if expected != getattr(crime, field):
            return False
    return True


def _matches_defendant_name(excel_name: str, json_name: str) -> bool:
    """遮蔽姓名依 ○ 前的中英文字配對、○ 數量不得大於實際長度；未遮蔽姓名維持完整比對。"""
    excel_name = _text(excel_name)
    json_name = _text(json_name)
    if "○" not in excel_name:
        return bool(excel_name) and excel_name == json_name
    excel_prefix, separator, excel_suffix = excel_name.partition("○")
    excel_suffix = separator + excel_suffix
    if not re.fullmatch(r"[A-Za-z\u3400-\u9fff\U00020000-\U0002FA1F]+", excel_prefix):
        return False
    json_suffix = json_name.split(excel_prefix, 1)[1] if excel_prefix in json_name else ""
    if len(excel_suffix) > len(json_suffix):
        return False
    return json_name.startswith(excel_prefix)


def _excel_col_to_val(column: str, value):
    text = _text(value)
    if column in {"未遂", "幫助", "累犯"}:
        if text in {"", "0", "0.0", "否", "N", "false"}:
            return False
        if text in {"1", "1.0", "是", "Y", "true", "✓"}:
            return True
        raise ValueError(f"{column}布林值無法辨識：{text}")
    if not text or text.lower() == "x":
        return None
    if column in {"案由", "裁判情形"}:
        return text
    if column in {"條段", "項段"}:
        return text.upper()
    number = _integer(value)
    if column in {"徒刑(YYMM)", "拘役", "罰金", "併科罰金"} and number == 0:
        return None
    if column == "徒刑(YYMM)":
        years, months = divmod(number, 100)
        if months >= 12:
            raise ValueError(f"YYMM 月份不合法：{value}")
        return years * 12 + months
    return number


def _load_analyses(path: Path, errors: list) -> dict[CaseNumber, list[tuple[Path, JudgmentAnalysis]]]:
    index = {}
    for source in _files(path, {".json"}):
        try:
            data = json.loads(source.read_text(encoding="utf-8-sig"))
            if not isinstance(data, dict) or "case_number" not in data:
                continue
            analysis = JudgmentAnalysis.from_dict(data)
            case_number = CaseNumber.from_string(analysis.case_number)
            index.setdefault(case_number, []).append((source, analysis))
        except (ValueError, TypeError, KeyError, OSError) as exc:
            errors.append({"file": str(source), "error": f"{type(exc).__name__}: {exc}"})
    return index


def _files(path: Path, suffixes: set[str]) -> list[Path]:
    if not path.exists():
        raise ValueError(f"路徑不存在：{path}")
    return sorted(p for p in ([path] if path.is_file() else path.rglob("*"))
                  if p.is_file() and p.suffix.lower() in suffixes and not p.name.startswith("~$"))


def _integer(value) -> int:
    number = float(value)
    if not number.is_integer() or number < 0:
        raise ValueError(f"非負整數格式錯誤：{value}")
    return int(number)


def _text(value) -> str:
    return "" if _blank(value) else re.sub(
        r"\s+", "", str(value)).replace("（", "(").replace("）", ")"
    )


def _blank(value) -> bool:
    return value is None or bool(pd.isna(value)) or str(value).strip() == ""


if __name__ == "__main__":
    main()
