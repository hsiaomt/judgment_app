from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from judgment_app.analysis.models import JudgmentAnalysis
from judgment_app.comparison.judgment import (
    build_judgment_analysis, compare_analyses, compare_files,
    excel_row_to_law_reference,
)
from judgment_app.comparison.common import compare_to_report, load_comparison_report


class CompareJudgmentTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.excel = self.root / "執行明細.xlsx"
        self.source = self.root / "判決.analysis.json"
        self.case = "臺灣高等法院115年度上訴字第1299號"
        self.row = {
            "序號": 22598267, "個案序號": 1,
            "裁判機關": "臺灣高院", "裁判年度": 115, "裁判冠字": "上訴",
            "裁判號數": "001299", "姓名": "陳○○", "累犯": None, "上訴類別": "01",
            "應執緩刑期間": 300, "緩刑法條": "01", "應執裁判情形": "科刑",
            "應執徒刑": 10500, "應執徒刑易科標準": "易科一仟折算",
            "應執拘役": 0, "應執拘役易科標準": None, "應執罰金": 0, "應執罰金易勞標準": None,
            "個案裁判情形": "科刑", "次數": 2, "個案案由": "詐欺", "法條": "3394x1x02",
            "個案徒刑": 10200, "個案徒刑易科標準": "易科一仟折算",
            "個案拘役": 0, "個案拘役易科標準": None, "個案罰金": 0, "個案罰金易勞標準": None,
        }
        self.data = {
            "case_number": self.case,
            "defendants": [{
                "name": "陳小明", "is_recidivist": False, "is_pros_appeal": True,
                "probation": 3, "probation_law": 1,
                "crimes": [{"crime_id": 9, "crime_name": "詐欺", "crime_jsst": "科刑",
                            "times": 2, "penalty": 14, "penalty_to_fine": 1000,
                            "laws": [{"law_name": "刑法", "article": 339, "article_sub": 4,
                                      "paragraph": 1, "paragraph_sub": 2,
                                      "is_attempted": False, "is_accessory": False}]}],
                "execution_groups": [{"crime_ids": [9], "executed_jsst": "科刑",
                                      "executed_penalty": 17, "executed_penalty_to_fine": 1000}],
            }],
        }

    def test_auto_comparison_selects_formats_and_writes_results(self):
        from judgment_app.comparison.common import compare_files_auto

        pd.DataFrame([self.row]).to_excel(self.excel, index=False, startrow=2)
        self.source.write_text(json.dumps(self.data), encoding="utf-8")
        law = {key: self.row[key] for key in
               ("裁判機關", "裁判年度", "裁判冠字", "裁判號數", "姓名")}
        law.update({"案由": "刑法", "條": 339, "之": 4, "項": 1, "款": 2,
                    "裁判情形": "科刑", "徒刑(YYMM)": 102})
        law_path = self.root / "檢核.xlsx"
        pd.DataFrame([law]).to_excel(law_path, index=False)
        messages = []
        report = compare_files_auto(self.root, self.source, progress=messages.append)
        self.assertEqual(report["summary"], {"equal": 2, "total": 2, "errors": 0})
        self.assertEqual(len(report["excel_outputs"]), 2)
        self.assertTrue(any("judgment" in message for message in messages))
        self.assertTrue(any("law_penalty" in message for message in messages))
        for path, header in ((self.excel, 3), (law_path, 1)):
            book = load_workbook(path)
            try:
                columns = {cell.value: cell.column for cell in book.active[header]}
                self.assertEqual(book.active.cell(header + 1, columns["status"]).value, "equal")
            finally:
                book.close()

    def test_auto_comparison_rejects_incomplete_unknown_and_mixed_formats(self):
        from judgment_app.comparison.common import compare_files_auto

        identity = {key: self.row[key] for key in
                    ("裁判機關", "裁判年度", "裁判冠字", "裁判號數", "姓名")}
        incomplete = dict(self.row)
        incomplete.pop("法條")
        for row in (identity, incomplete):
            with self.subTest(columns=list(row)):
                pd.DataFrame([row]).to_excel(self.excel, index=False)
                before = self.excel.read_bytes()
                report = compare_files_auto(self.excel, self.root)
                self.assertEqual(report["rows"], [])
                self.assertEqual(report["errors"][0]["stage"], "detect_comparison")
                self.assertEqual(self.excel.read_bytes(), before)
        with pd.ExcelWriter(self.excel) as writer:
            pd.DataFrame([self.row]).to_excel(writer, sheet_name="明細", index=False)
            pd.DataFrame([{**identity, "案由": "刑法", "條": 339, "裁判情形": "科刑"}]).to_excel(
                writer, sheet_name="檢核", index=False)
        report = compare_files_auto(self.excel, self.root)
        self.assertIn("混用", report["errors"][0]["error"])
        self.assertEqual(report["excel_outputs"], [])

    def run_report(self, rows=None):
        pd.DataFrame(rows or [self.row]).to_excel(self.excel, index=False, startrow=2)
        self.source.write_text(json.dumps(self.data), encoding="utf-8")
        report = compare_files(self.excel, self.root)
        if report['excel_outputs']:
            book = load_workbook(self.excel)
            try:
                sheet = book.active
                columns = {cell.value: cell.column for cell in sheet[3]}
                for result in report['rows']:
                    for location in result.get('excel_rows', [result]):
                        self.assertEqual(sheet.cell(location['row'], columns['status']).value, result['status'])
                        self.assertEqual(sheet.cell(location['row'], columns['reason']).value, result.get('reason'))
            finally:
                book.close()
        return report

    def test_full_comparison_and_masked_name(self):
        report = self.run_report()
        self.assertEqual(report["summary"], {"equal": 1, "total": 1, "errors": 0})
        self.assertEqual(report["rows"][0]["excel_rows"][0]["row"], 4)
        self.assertEqual(report["sheets"][0]["defaulted_fields"], ["is_accessory", "is_attempted"])

    def test_same_crime_sequence_merges_laws_without_multiplying_times(self):
        second = {**self.row, "法條": "0190x1B00"}
        law = {"law_name": "刑法", "article": 19, "paragraph": 1, "paragraph_part": "B",
               "is_attempted": False, "is_accessory": False}
        self.data["defendants"][0]["crimes"][0]["laws"].append(law)
        analysis = build_judgment_analysis(self.case, [self.row, second, self.row])
        defendant = analysis.defendants[0]
        self.assertEqual(len(defendant.crimes), 1)
        self.assertEqual(defendant.crimes[0].times, 2)
        self.assertEqual(len(defendant.crimes[0].laws), 2)
        self.assertEqual(defendant.execution_groups[0].crime_ids, [1])
        self.assertEqual(self.run_report([second, self.row])["rows"][0]["status"], "equal")

    def test_same_sequence_different_crime_sequences_are_one_group(self):
        defendant = self.data["defendants"][0]
        defendant["crimes"].append({**deepcopy(defendant["crimes"][0]), "crime_id": 10})
        defendant["execution_groups"][0]["crime_ids"].append(10)
        rows = [self.row, {**self.row, "個案序號": 2}]
        analysis = build_judgment_analysis(self.case, rows)
        self.assertEqual(len(analysis.defendants[0].crimes), 2)
        self.assertEqual(analysis.defendants[0].execution_groups[0].crime_ids, [1, 2])
        self.assertEqual(self.run_report(rows)["rows"][0]["status"], "equal")

    def test_different_sequences_same_name_are_distinct_defendants(self):
        rows = [self.row, {**self.row, "序號": 22598268, "次數": 3}]
        analysis = build_judgment_analysis(self.case, rows)
        self.assertEqual(len(analysis.defendants), 2)
        self.assertEqual([d.name for d in analysis.defendants], ["陳○○", "陳○○"])
        self.assertEqual([len(d.crimes) for d in analysis.defendants], [1, 1])
        self.assertEqual([d.crimes[0].times for d in analysis.defendants], [2, 3])
        self.assertEqual([d.execution_groups[0].crime_ids for d in analysis.defendants], [[1], [1]])
        report = self.run_report(rows)
        self.assertEqual([r["status"] for r in report["rows"]], ["equal", "unmatched"])

    def test_same_sequence_conflicting_names_are_invalid(self):
        rows = [self.row, {**self.row, "姓名": "李○○"}]
        with self.assertRaisesRegex(ValueError, "同一序號的被告姓名不一致"):
            build_judgment_analysis(self.case, rows)
        report = self.run_report(rows)
        self.assertEqual(report["summary"], {"unmatched": 1, "total": 1, "errors": 1})
        self.assertEqual(len(report["rows"][0]["excel_rows"]), 2)

    def test_each_execution_rate_separates_groups(self):
        for column, field in (("應執徒刑易科標準", "executed_penalty_to_fine"),
                              ("應執拘役易科標準", "executed_detention_to_fine"),
                              ("應執罰金易勞標準", "executed_fine_to_detention")):
            with self.subTest(column=column):
                rows = [self.row, {**self.row, "個案序號": 2, column: 2000, "應執徒刑": 10600}]
                analysis = build_judgment_analysis(self.case, rows)
                defendant = analysis.defendants[0]
                self.assertEqual(len(defendant.execution_groups), 2)
                self.assertEqual([g.crime_ids for g in defendant.execution_groups], [[1], [2]])
                self.assertEqual(getattr(defendant.execution_groups[1], field), 2000)
                self.assertEqual(defendant.execution_groups[1].executed_penalty, 18)
                self.data = analysis.to_dict()
                self.assertEqual(self.run_report(rows)["rows"][0]["status"], "equal")
                self.data["defendants"][0]["execution_groups"].pop()
                self.assertEqual(self.run_report(rows)["rows"][0]["status"], "unmatched")

    def test_equivalent_execution_rates_share_group(self):
        rows = [self.row, {**self.row, "個案序號": 2, "應執徒刑易科標準": 1000}]
        defendant = build_judgment_analysis(self.case, rows).defendants[0]
        self.assertEqual(len(defendant.execution_groups), 1)
        self.assertEqual(defendant.execution_groups[0].crime_ids, [1, 2])

    def test_same_crime_is_referenced_by_multiple_groups(self):
        rows = [self.row, {**self.row, "法條": "0190x1B00", "應執徒刑易科標準": 2000}]
        defendant = build_judgment_analysis(self.case, rows).defendants[0]
        self.assertEqual(len(defendant.crimes), 1)
        self.assertEqual(len(defendant.crimes[0].laws), 2)
        self.assertEqual([g.crime_ids for g in defendant.execution_groups], [[1], [1]])

    def test_every_ignored_field(self):
        self.data.update(lower_court_case_number="前審", pros_appearing=["檢察官"])
        defendant = self.data["defendants"][0]
        defendant.update(confiscation=["物品"], appeal_result="駁回")
        defendant["crimes"][0].update(crime_id=42, crime_name="不同罪名")
        defendant["execution_groups"][0].update(crime_ids=[99])
        self.row.update(蒞庭檢察官="忽略", 上訴結果="忽略", 個案併科罰金="忽略")
        self.assertEqual(self.run_report()["rows"][0]["status"], "equal")

    def test_nonignored_fields_are_compared(self):
        mutations = [
            ((), "case_number", "臺灣高等法院115年度上訴字第1300號"),
            (("defendants", 0), "is_recidivist", True),
            (("defendants", 0), "is_pros_appeal", False),
            (("defendants", 0), "probation", 2),
            (("defendants", 0), "probation_law", 2),
        ]
        mutations += [(('defendants', 0, 'crimes', 0), field, value) for field, value in (
            ("crime_jsst", "免刑"), ("times", 3), ("penalty", 15), ("penalty_to_fine", 2000),
            ("detention", 10), ("detention_to_fine", 1000), ("fine", 10000), ("fine_to_detention", 1000),
        )]
        mutations += [(('defendants', 0, 'execution_groups', 0), field, value) for field, value in (
            ("executed_jsst", "免刑"), ("executed_penalty", 18), ("executed_penalty_to_fine", 2000),
            ("executed_detention", 10), ("executed_detention_to_fine", 1000),
            ("executed_fine", 10000), ("executed_fine_to_detention", 1000),
        )]
        mutations += [(('defendants', 0, 'crimes', 0, 'laws', 0), field, value) for field, value in (
            ("law_name", "洗錢防制法"), ("article", 340), ("article_sub", 5), ("article_part", "A"),
            ("paragraph", 2), ("paragraph_part", "B"), ("paragraph_sub", 3),
            ("is_attempted", True), ("is_accessory", True),
        )]
        expected = build_judgment_analysis(self.case, [self.row])
        for path, field, value in mutations:
            with self.subTest(field=field):
                data = deepcopy(self.data)
                target = data
                for key in path:
                    target = target[key]
                target[field] = value
                result = compare_analyses(expected, JudgmentAnalysis.from_dict(data))
                self.assertEqual(result["status"], "unmatched")
                self.assertTrue(result["differences"])
                self.assertIn(field, result["reason"])

    def test_specific_reasons_are_written_to_excel(self):
        self.data['defendants'][0]['probation'] = 4
        self.data['defendants'][0]['crimes'][0]['times'] = 5
        self.data['defendants'][0]['crimes'][0]['laws'][0]['article'] = 340
        report = self.run_report()
        reason = report['rows'][0]['reason']
        for field in ('defendants[0].probation', 'defendants[0].crimes[].times',
                      'defendants[0].crimes[].laws[].article'):
            self.assertIn(f'{field} 不合', reason)
        self.assertNotIn('JudgmentAnalysis 內容不合', reason)

    def test_crime_length_is_ignored_but_fields_must_match_one_candidate(self):
        expected = build_judgment_analysis(self.case, [self.row])
        actual = expected.model_copy(deep=True)
        actual.defendants[0].crimes.append(actual.defendants[0].crimes[0].model_copy(deep=True))
        self.assertEqual(compare_analyses(expected, actual)['status'], 'equal')
        expected = actual.model_copy(deep=True)
        expected.defendants[0].crimes[0].times = 3
        expected.defendants[0].crimes[1].penalty = 20
        actual = expected.model_copy(deep=True)
        actual.defendants[0].crimes[0].times = 2
        actual.defendants[0].crimes[1].times = 3
        result = compare_analyses(expected, actual)
        self.assertEqual(result['status'], 'unmatched')
        self.assertNotIn('crimes.length', result['reason'])

    def test_extra_and_missing_items_are_detected(self):
        expected = build_judgment_analysis(self.case, [self.row])
        for path in (("defendants",), ("defendants", 0, "crimes"),
                     ("defendants", 0, "execution_groups"), ("defendants", 0, "crimes", 0, "laws")):
            for extra in (True, False):
                with self.subTest(path=path, extra=extra):
                    data = deepcopy(self.data)
                    target = data
                    for key in path:
                        target = target[key]
                    if extra:
                        item = deepcopy(target[0])
                        if path == ("defendants",):
                            item["name"] = "李大華"
                        target.append(item)
                    else:
                        target.clear()
                    result = compare_analyses(expected, JudgmentAnalysis.from_dict(data))
                    expected_status = 'equal' if path in (('defendants',), ('defendants', 0, 'execution_groups'),
                                                        ('defendants', 0, 'crimes')) and extra else 'unmatched'
                    self.assertEqual(result["status"], expected_status)

    def test_order_does_not_matter_and_inputs_are_unchanged(self):
        data = deepcopy(self.data)
        defendant = data["defendants"][0]
        defendant["crimes"].append({**deepcopy(defendant["crimes"][0]), "crime_id": 10, "times": 3})
        defendant["execution_groups"].append({**defendant["execution_groups"][0], "crime_ids": [10], "executed_penalty": 20})
        defendant["crimes"][0]["laws"].append({**defendant["crimes"][0]["laws"][0], "article": 340})
        data["defendants"].append({**deepcopy(defendant), "name": "李大華"})
        expected = JudgmentAnalysis.from_dict(data)
        for defendant in data["defendants"]:
            for crime in defendant["crimes"]:
                crime["laws"].reverse()
            defendant["crimes"].reverse()
            defendant["execution_groups"].reverse()
        data["defendants"].reverse()
        actual = JudgmentAnalysis.from_dict(data)
        before = actual.model_dump()
        self.assertEqual(compare_analyses(expected, actual)["status"], "equal")
        self.assertEqual(actual.model_dump(), before)

    def test_more_excel_crimes_can_reuse_one_matching_json_crime(self):
        report = self.run_report([self.row, {**self.row, '個案序號': 2}])
        self.assertEqual(report['rows'][0]['status'], 'equal')
        self.assertNotIn('reason', report['rows'][0])

    def test_extra_different_json_crime_does_not_affect_matching(self):
        extra = deepcopy(self.data['defendants'][0]['crimes'][0])
        extra.update(crime_id=10, penalty=99, times=20)
        self.data['defendants'][0]['crimes'].insert(0, extra)
        self.assertEqual(self.run_report()['rows'][0]['status'], 'equal')

    def test_unmatched_excel_crime_still_fails_when_lengths_differ(self):
        report = self.run_report([self.row, {**self.row, '個案序號': 2, '次數': 3}])
        self.assertEqual(report['rows'][0]['status'], 'unmatched')
        self.assertIn('crimes[].times', report['rows'][0]['reason'])
        self.assertNotIn('crimes.length', report['rows'][0]['reason'])

    def test_any_complete_matching_defendant_passes_regardless_of_order(self):
        other = deepcopy(self.data["defendants"][0])
        other.update(name="陳大華", probation=5)
        self.data["defendants"].insert(0, other)
        self.assertEqual(self.run_report()["rows"][0]["status"], "equal")
        self.data["defendants"].reverse()
        self.assertEqual(self.run_report()["rows"][0]["status"], "equal")
        self.data["defendants"].pop()
        report = self.run_report([self.row, {**self.row, "序號": 22598268, "姓名": "陳小○"}])
        self.assertEqual([result['status'] for result in report['rows']], ['equal', 'equal'])

    def test_multiple_candidates_cannot_combine_fields_and_report_closest(self):
        other = deepcopy(self.data['defendants'][0])
        other.update(name='陳大華', probation=5)
        self.data['defendants'][0]['crimes'][0]['times'] = 3
        self.data['defendants'][0]['execution_groups'][0]['executed_penalty'] = 18
        self.data['defendants'].append(other)
        expected = build_judgment_analysis(self.case, [self.row])
        actual = JudgmentAnalysis.from_dict(self.data)
        before_expected, before_actual = expected.model_copy(deep=True), actual.model_copy(deep=True)
        result = compare_analyses(expected, actual)
        self.assertEqual(result['status'], 'unmatched')
        self.assertEqual(result['reason'], 'defendants[0].probation 不合')
        self.assertEqual(expected, before_expected)
        self.assertEqual(actual, before_actual)
        self.assertEqual(self.run_report()['rows'][0]['reason'], result['reason'])

    def test_matching_fields_with_nonmatching_name_do_not_pass(self):
        other = deepcopy(self.data['defendants'][0])
        other['name'] = '李大華'
        self.data['defendants'][0]['probation'] = 5
        self.data['defendants'].append(other)
        self.assertEqual(self.run_report()['rows'][0]['status'], 'unmatched')

    def test_conflicts_fail_defendant_with_sequence_context(self):
        for change in ({"累犯": 1}, {"個案徒刑": 10300}, {"次數": 3},
                       {"應執徒刑": 10600}, {"應執緩刑期間": 400}, {"上訴類別": None}):
            with self.subTest(change=change):
                report = self.run_report([self.row, {**self.row, **change}])
                self.assertEqual(report["rows"][0]["status"], "unmatched")
                self.assertIn("序號=22598267", report["errors"][0]["error"])
                self.assertIn("不一致", report["errors"][0]["error"])

    def test_only_different_defendant_is_unmatched_in_report_and_excel(self):
        other = deepcopy(self.data['defendants'][0])
        other.update(name='李大華', probation=4)
        self.data['defendants'].append(other)
        report = self.run_report([self.row, {**self.row, '序號': 22598268, '姓名': '李○○'}])
        self.assertEqual(report['summary'], {'equal': 1, 'unmatched': 1, 'total': 2, 'errors': 0})
        first, second = report['rows']
        self.assertEqual((first['name'], first['status']), ('陳○○', 'equal'))
        self.assertNotIn('reason', first)
        self.assertEqual(first['excel_rows'][0]['row'], 4)
        self.assertEqual((second['name'], second['status']), ('李○○', 'unmatched'))
        self.assertIn('probation', second['reason'])
        self.assertEqual(second['excel_rows'][0]['row'], 5)

    def test_invalid_excel_defendant_does_not_block_other_defendant(self):
        other = {**deepcopy(self.data['defendants'][0]), 'name': '李大華'}
        self.data['defendants'].append(other)
        report = self.run_report([{**self.row, '序號': 22598268, '姓名': '李○○', '次數': 0}, self.row])
        self.assertEqual([r['status'] for r in report['rows']], ['unmatched', 'equal'])
        self.assertEqual(report['summary']['errors'], 1)
        self.assertEqual(report['errors'][0]['name'], '李○○')

    def test_multiple_or_missing_candidates_do_not_block_other_defendant(self):
        self.data['defendants'].append({**deepcopy(self.data['defendants'][0]), 'name': '李大華'})
        for ambiguous in (False, True):
            with self.subTest(ambiguous=ambiguous):
                if ambiguous:
                    self.data['defendants'].append({**deepcopy(self.data['defendants'][0]), 'name': '陳大華'})
                name = '陳○○' if ambiguous else '王○○'
                report = self.run_report([{**self.row, '姓名': name}, {**self.row, '序號': 22598268, '姓名': '李○○'}])
                self.assertEqual([r['status'] for r in report['rows']],
                                 ['equal' if ambiguous else 'unmatched', 'equal'])
                if not ambiguous:
                    self.assertIn('找不到姓名相符的被告', report['rows'][0]['reason'])

    def test_case_number_is_checked_before_defendant_name(self):
        expected = build_judgment_analysis(self.case, [self.row])
        actual = JudgmentAnalysis.from_dict({**self.data, 'case_number': '臺灣高等法院115年度上訴字第1號',
                                            'defendants': []})
        with patch('judgment_app.comparison.judgment._compare_defendant',
                   side_effect=AssertionError('案號不同不應比對被告')):
            result = compare_analyses(expected, actual)
        self.assertEqual(result['reason'], 'case_number 不合')
        self.assertEqual(len(result['defendant_results']), 1)
        self.assertEqual(result['defendant_results'][0]['reason'], 'case_number 不合')

    def test_compare_analyses_returns_all_defendants_after_missing_name(self):
        expected = build_judgment_analysis(self.case, [
            {**self.row, '姓名': '王○○'},
            {**self.row, '序號': 22598268},
            {**self.row, '序號': 22598269, '應執緩刑期間': 400},
        ])
        result = compare_analyses(expected, JudgmentAnalysis.from_dict(self.data))
        self.assertEqual(result['status'], 'unmatched')
        self.assertIn('找不到姓名相符的被告', result['reason'])
        self.assertEqual([r['status'] for r in result['defendant_results']],
                         ['unmatched', 'equal', 'unmatched'])
        self.assertEqual(result['defendant_results'][2]['reason'], 'defendants[2].probation 不合')

    def test_each_defendant_can_match_a_different_document(self):
        other = deepcopy(self.data)
        other['defendants'][0]['name'] = '李大華'
        (self.root / '另一份.json').write_text(json.dumps(other), encoding='utf-8')
        compared = []

        def capture(expected, actual):
            result = compare_analyses(expected, actual)
            compared.append((expected, actual, result))
            return result

        with patch('judgment_app.comparison.judgment.compare_analyses', side_effect=capture) as compare:
            report = self.run_report([self.row, {**self.row, '序號': 22598268, '姓名': '李○○'}])
        self.assertEqual(compare.call_count, 2)
        self.assertIs(compared[0][0], compared[1][0])
        self.assertEqual(len(compared[0][0].defendants), 2)
        for _, actual, result in compared:
            self.assertEqual(result['status'], 'unmatched')
            expected_statuses = ['equal', 'unmatched'] if actual.defendants[0].name == '陳小明' else ['unmatched', 'equal']
            self.assertEqual([r['status'] for r in result['defendant_results']], expected_statuses)
        self.assertEqual(report['summary'], {'equal': 2, 'total': 2, 'errors': 0})
        for result in report['rows']:
            self.assertEqual({r['status'] for r in result['document_results']}, {'equal', 'unmatched'})

    def test_missing_ids_and_zero_times_are_invalid(self):
        for change in ({"序號": None}, {"個案序號": None}, {"次數": 0}):
            with self.subTest(change=change):
                report = self.run_report([{**self.row, **change}])
                self.assertEqual(report["summary"]["errors"], 1)
                self.assertEqual(report["rows"][0]["status"], "unmatched")

    def test_appeal_type_any_value_is_true_blank_depends_on_instance(self):
        for value in ("01", "02", "任意非空", 0, 1):
            with self.subTest(value=value):
                result = build_judgment_analysis(self.case, [{**self.row, "上訴類別": value}])
                self.assertIs(result.defendants[0].is_pros_appeal, True)
        for case, expected in ((self.case, False), ("臺灣臺北地方法院115年度審訴字第1號", None),
                               ("臺灣臺北地方法院115年度簡上字第1號", False)):
            with self.subTest(case=case):
                result = build_judgment_analysis(case, [{**self.row, "上訴類別": None}])
                self.assertIs(result.defendants[0].is_pros_appeal, expected)

    def test_sentence_units_rates_and_invalid_months_days(self):
        for raw, expected in ((0, None), (300, 3), (10200, 14), (300000, 360)):
            result = build_judgment_analysis(self.case, [{**self.row, "個案徒刑": raw}])
            self.assertEqual(result.defendants[0].crimes[0].penalty, expected)
        for raw in (1200, 101, -1):
            with self.subTest(raw=raw), self.assertRaises(ValueError):
                build_judgment_analysis(self.case, [{**self.row, "個案徒刑": raw}])
        for raw, expected in (("易科二仟折算", 2000), ("易科三仟折算", 3000), (1000, 1000)):
            result = build_judgment_analysis(self.case, [{**self.row, "個案徒刑易科標準": raw}])
            self.assertEqual(result.defendants[0].crimes[0].penalty_to_fine, expected)
        row = {**self.row, "個案罰金": 3000, "個案罰金易勞標準": "易勞一仟折算"}
        result = build_judgment_analysis(self.case, [row])
        self.assertEqual(result.defendants[0].crimes[0].fine, 3000)
        self.assertEqual(result.defendants[0].crimes[0].fine_to_detention, 1000)

    def test_law_code_parts_aliases_and_invalid_codes(self):
        for code, field, expected in (("1051x1A00", "paragraph_part", "A"),
                                      ("0430A0x00", "article_part", "A"),
                                      ("3394x1x02", "paragraph_sub", 2)):
            law = excel_row_to_law_reference({**self.row, "法條": code})
            self.assertEqual(getattr(law, field), expected)
        law = excel_row_to_law_reference({**self.row, "個案案由": "毒品防制條例"})
        self.assertEqual(law.law_name, "毒品危害防制條例")
        for code in ("000000000", 0, None):
            self.assertIsNone(excel_row_to_law_reference({**self.row, "法條": code}))
        for code in ("3394x1x2", "3394Q1x02", "bad"):
            with self.subTest(code=code), self.assertRaises(ValueError):
                excel_row_to_law_reference({**self.row, "法條": code})

    def test_case_rows_span_sheets_and_files(self):
        other = {**self.row, "序號": 22598268, "姓名": "李○○"}
        self.data["defendants"].append({**deepcopy(self.data["defendants"][0]), "name": "李大華", "probation": 4})
        self.source.write_text(json.dumps(self.data), encoding="utf-8")
        with pd.ExcelWriter(self.excel) as writer:
            pd.DataFrame([self.row]).to_excel(writer, sheet_name="一", index=False)
            pd.DataFrame([other]).to_excel(writer, sheet_name="二", index=False)
        report = compare_files(self.excel, self.root)
        self.assertEqual([r['status'] for r in report['rows']], ['equal', 'unmatched'])
        book = load_workbook(self.excel)
        for sheet_name, result in zip(('一', '二'), report['rows']):
            sheet = book[sheet_name]
            columns = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, columns['status']).value, result['status'])
            self.assertEqual(sheet.cell(2, columns['reason']).value, result.get('reason'))
        book.close()
        pd.DataFrame([self.row]).to_excel(self.excel, index=False)
        pd.DataFrame([other]).to_excel(self.root / "另一份.xlsx", index=False)
        report = compare_files(self.root, self.root)
        self.assertEqual(report['summary'], {'equal': 1, 'unmatched': 1, 'total': 2, 'errors': 0})
        for result in report['rows']:
            book = load_workbook(result['excel_rows'][0]['excel'])
            sheet = book.active
            columns = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, columns['status']).value, result['status'])
            self.assertEqual(sheet.cell(2, columns['reason']).value, result.get('reason'))
            book.close()

    def test_different_documents_are_not_combined(self):
        other = deepcopy(self.data)
        other["defendants"][0]["probation"] = 5
        (self.root / "另一份.json").write_text(json.dumps(other), encoding="utf-8")
        result = self.run_report()["rows"][0]
        self.assertEqual(result["status"], "equal")
        self.assertEqual({r["status"] for r in result["document_results"]}, {"equal", "unmatched"})
        self.data["defendants"][0]["crimes"][0]["times"] = 5
        self.assertEqual(self.run_report()["rows"][0]["status"], "unmatched")

    def test_nonguilty_and_exempt_crimes_are_compared(self):
        for status in ("無罪", "免訴", "不受理", "免刑"):
            with self.subTest(status=status):
                row = {**self.row, "個案裁判情形": status, "應執裁判情形": status,
                       "個案徒刑": 0, "應執徒刑": 0, "個案徒刑易科標準": None,
                       "應執徒刑易科標準": None, "法條": "000000000"}
                expected = build_judgment_analysis(self.case, [row])
                actual = expected.model_copy(deep=True)
                actual.defendants[0].name = "陳小明"
                self.assertEqual(compare_analyses(expected, actual)["status"], "equal")

    def test_execution_status_is_used_independently_of_crime_status(self):
        for status in ('無罪', '免訴', '不受理', '免刑'):
            with self.subTest(status=status):
                self.row['應執裁判情形'] = status
                expected = build_judgment_analysis(self.case, [self.row])
                self.assertEqual(expected.defendants[0].crimes[0].crime_jsst, '科刑')
                self.assertEqual(expected.defendants[0].execution_groups[0].executed_jsst, status)
                self.assertEqual(self.run_report()['rows'][0]['status'], 'unmatched')
        for value in (None, '不合法內容'):
            with self.subTest(value=value):
                self.row['應執裁判情形'] = value
                self.assertEqual(self.run_report()['summary']['errors'], 1)
        del self.row['應執裁判情形']
        report = self.run_report()
        self.assertTrue(any('應執裁判情形' in e['error'] for e in report['errors']))

    def test_mixed_execution_statuses_in_one_sequence_find_separate_groups(self):
        rows = [{**self.row, '個案裁判情形': status, '個案序號': index,
                 '應執裁判情形': status, '個案徒刑': 0, '個案徒刑易科標準': None,
                 '法條': '000000000'}
                for index, status in enumerate(('無罪', '免訴', '不受理', '免刑'), 1)]
        expected = build_judgment_analysis(self.case, rows)
        actual = expected.model_copy(deep=True)
        for group in actual.defendants[0].execution_groups:
            group.executed_penalty = None
            group.executed_penalty_to_fine = None
        self.assertEqual(compare_analyses(expected, actual)['status'], 'equal')
        for index in range(4):
            with self.subTest(index=index):
                incomplete = actual.model_copy(deep=True)
                incomplete.defendants[0].execution_groups.pop(index)
                self.assertIn('executed_jsst', compare_analyses(expected, incomplete)['reason'])

    def test_multiple_excel_groups_can_reuse_one_json_group_for_every_status(self):
        for status in ('科刑', '無罪', '免訴', '不受理', '免刑'):
            with self.subTest(status=status):
                expected = build_judgment_analysis(
                    self.case, [{**self.row, '個案裁判情形': status, '應執裁判情形': status}],
                )
                expected.defendants[0].execution_groups.append(
                    expected.defendants[0].execution_groups[0].model_copy(deep=True),
                )
                actual = expected.model_copy(deep=True)
                actual.defendants[0].execution_groups.pop()
                before = actual.model_dump()
                self.assertEqual(compare_analyses(expected, actual)['status'], 'equal')
                self.assertEqual(actual.model_dump(), before)
                actual.defendants[0].execution_groups[0].executed_jsst = '無罪' if status != '無罪' else '科刑'
                self.assertEqual(compare_analyses(expected, actual)['status'], 'unmatched')

    def test_case_is_built_once_with_all_defendants_across_sheets(self):
        self.source.write_text(json.dumps(self.data), encoding='utf-8')
        other = {**self.row, '序號': 22598268, '姓名': '李○○'}
        with pd.ExcelWriter(self.excel) as writer:
            pd.DataFrame([self.row, other]).to_excel(writer, sheet_name='一', index=False)
            pd.DataFrame([{**self.row, '法條': '0190x1B00'}]).to_excel(writer, sheet_name='二', index=False)
        built = []

        def capture(*args, **kwargs):
            analysis = build_judgment_analysis(*args, **kwargs)
            built.append(analysis)
            return analysis

        with patch('judgment_app.comparison.judgment.build_judgment_analysis', side_effect=capture) as builder:
            report = compare_files(self.excel, self.root)
        self.assertEqual(builder.call_count, 1)
        self.assertEqual(len(built[0].defendants), 2)
        self.assertEqual(len(built[0].defendants[0].crimes[0].laws), 2)
        self.assertEqual(len(report['rows'][0]['excel_rows']), 2)

    def test_comparison_and_building_do_not_serialize_models_or_json(self):
        actual = JudgmentAnalysis.from_dict(self.data)
        with patch('judgment_app.analysis.models.AnalysisModel.model_dump',
                   side_effect=AssertionError('比較不應序列化模型')), patch(
            'judgment_app.comparison.judgment.json.dumps',
            side_effect=AssertionError('比較不應序列化 JSON'),
        ):
            expected = build_judgment_analysis(self.case, [self.row, self.row])
            self.assertEqual(compare_analyses(expected, actual)['status'], 'equal')
            actual.defendants[0].crimes[0].laws[0] = actual.defendants[0].crimes[0].laws[0].model_copy(
                update={'article': 340},
            )
            result = compare_analyses(expected, actual)
            self.assertEqual(result['status'], 'unmatched')
            self.assertIn('laws[].article', result['reason'])

    def test_invalid_case_json_and_missing_headers(self):
        report = self.run_report([self.row, {**self.row, "裁判年度": "無效"}])
        self.assertEqual(report["summary"]["errors"], 1)
        self.assertEqual(report["summary"]["unmatched"], 1)
        del self.data["defendants"][0]["crimes"][0]["laws"]
        report = self.run_report()
        self.assertEqual(report["rows"][0]["reason"], "判決不存在")
        self.assertEqual(report["rows"][0]["status"], "judgment_not_found")
        self.assertEqual(report["summary"]["errors"], 1)
        del self.row["個案序號"]
        report = self.run_report()
        self.assertTrue(any("缺少執行明細欄位" in e["error"] for e in report["errors"]))

    def test_missing_judgment_has_separate_status_and_summary(self):
        report = self.run_report([self.row, {**self.row, "裁判號數": 1300}])
        self.assertEqual(report['summary'], {'equal': 1, 'judgment_not_found': 1, 'total': 2, 'errors': 0})
        self.assertEqual(report['rows'][1]['status'], 'judgment_not_found')
        self.assertEqual(report['rows'][1]['reason'], '判決不存在')

    def test_report_output_and_source_protection(self):
        self.run_report()
        destination = self.root / "report.json"
        report = compare_to_report(self.excel, self.root, destination, compare=compare_files)
        self.assertEqual(json.loads(destination.read_text(encoding="utf-8")), report)
        self.assertEqual(load_comparison_report(destination), report)
        with self.assertRaisesRegex(ValueError, "覆寫分析 JSON"):
            compare_to_report(self.excel, self.root, self.source, compare=compare_files)

    def test_grouped_mismatch_report_can_be_loaded(self):
        self.data['defendants'][0]['crimes'][0]['penalty'] = 12
        report = self.run_report([self.row, self.row])
        self.assertEqual(report['summary'], {'unmatched': 1, 'total': 1, 'errors': 0})
        result = report['rows'][0]
        self.assertEqual(len(result['excel_rows']), 2)
        self.assertTrue(result['document_results'][0]['differences'])
        destination = self.root / 'report.json'
        destination.write_text(json.dumps(report), encoding='utf-8')
        self.assertEqual(load_comparison_report(destination), report)

    def test_rerun_clears_reasons_and_empty_rows_without_extra_columns(self):
        self.row['姓名'] = '不存在'
        self.run_report()
        book = load_workbook(self.excel)
        sheet = book.active
        columns = {cell.value: cell.column for cell in sheet[3]}
        sheet.cell(4, columns['姓名'], '陳○○')
        sheet.cell(5, columns['status'], 'unmatched')
        sheet.cell(5, columns['reason'], '舊原因')
        book.save(self.excel)
        book.close()
        report = compare_files(self.excel, self.root)
        self.assertEqual(report['summary']['equal'], 1)
        self.assertEqual(report['summary']['total'], 1)
        self.assertNotIn('reason', report['sheets'][0]['uncompared_columns'])
        book = load_workbook(self.excel)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.max_column, len(self.row) + 2)
        self.assertEqual(sheet.cell(4, columns['status']).value, 'equal')
        self.assertIsNone(sheet.cell(4, columns['reason']).value)
        self.assertIsNone(sheet.cell(5, columns['status']).value)
        self.assertIsNone(sheet.cell(5, columns['reason']).value)

    def test_save_failure_preserves_workbook_and_reports_error(self):
        self.run_report()
        original = self.excel.read_bytes()
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=PermissionError('無法寫入')):
            report = compare_files(self.excel, self.root)
        self.assertEqual(self.excel.read_bytes(), original)
        self.assertEqual(report['excel_outputs'], [])
        self.assertEqual(report['errors'][0]['stage'], 'write_excel')
        self.assertEqual(report['summary']['errors'], 1)
        self.assertEqual(list(self.root.glob('~$status-*')), [])

    def test_cross_sheet_case_failure_is_written_to_every_source_row(self):
        self.source.write_text(json.dumps(self.data), encoding='utf-8')
        with pd.ExcelWriter(self.excel) as writer:
            pd.DataFrame([self.row]).to_excel(writer, sheet_name='一', index=False)
            pd.DataFrame([{**self.row, '次數': 3}]).to_excel(writer, sheet_name='二', index=False)
        report = compare_files(self.excel, self.root)
        result = report['rows'][0]
        self.assertEqual(result['status'], 'unmatched')
        book = load_workbook(self.excel)
        self.addCleanup(book.close)
        for sheet in book:
            columns = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, columns['status']).value, 'unmatched')
            self.assertEqual(sheet.cell(2, columns['reason']).value, result['reason'])

    def test_xls_export_preserves_all_sheets_and_does_not_overwrite(self):
        self.run_report()
        book = load_workbook(self.excel)
        book.create_sheet('說明')['A1'] = '保留此頁'
        book.save(self.excel)
        book.close()
        source = self.root / 'legacy.xls'
        source.write_bytes(b'legacy')
        excel_file = pd.ExcelFile
        with patch('judgment_app.comparison.judgment.pd.ExcelFile',
                   side_effect=lambda path: excel_file(self.excel if Path(path) == source else path)):
            report = compare_files(source, self.root)
            output = Path(report['excel_outputs'][0]['output'])
            self.assertEqual(output.name, 'legacy.status.xlsx')
            original = output.read_bytes()
            book = load_workbook(output)
            self.assertEqual(book['說明']['A1'].value, '保留此頁')
            sheet = book.worksheets[0]
            columns = {cell.value: cell.column for cell in sheet[3]}
            self.assertEqual(sheet.cell(4, columns['status']).value, 'equal')
            self.assertIsNone(sheet.cell(4, columns['reason']).value)
            book.close()
            report = compare_files(self.root, self.root)
            self.assertEqual(output.read_bytes(), original)
            self.assertTrue((self.root / 'legacy.status-2.xlsx').exists())
            self.assertEqual(len(report['rows'][0]['excel_rows']), 2)
        self.assertEqual(source.read_bytes(), b'legacy')

    def test_xls_read_failure_does_not_export_partial_workbook(self):
        self.run_report()
        book = load_workbook(self.excel)
        book.create_sheet('讀取失敗')
        book.save(self.excel)
        book.close()
        source = self.root / 'legacy.xls'
        source.write_bytes(b'legacy')
        excel_file, parse = pd.ExcelFile, pd.ExcelFile.parse

        def read_sheet(book, sheet, **kwargs):
            if sheet == '讀取失敗':
                raise ValueError('讀取失敗')
            return parse(book, sheet, **kwargs)

        with patch.object(excel_file, 'parse', read_sheet), patch(
            'judgment_app.comparison.judgment.pd.ExcelFile', side_effect=lambda path: excel_file(self.excel),
        ):
            report = compare_files(source, self.root)
        self.assertEqual(report['rows'][0]['status'], 'equal')
        self.assertEqual(report['excel_outputs'], [])
        self.assertTrue(any(e.get('stage') == 'write_excel' for e in report['errors']))
        self.assertFalse((self.root / 'legacy.status.xlsx').exists())


if __name__ == "__main__":
    unittest.main()
