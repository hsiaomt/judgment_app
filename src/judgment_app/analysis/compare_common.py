"""Excel 比對共用的欄位轉換、姓名配對及檔案讀寫工具。"""

import json
from collections.abc import Iterator
from pathlib import Path
import re
from tempfile import NamedTemporaryFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from judgment_app.analysis.analysis_models import JudgmentAnalysis
from judgment_app.analysis.comparison_models import ComparisonReport
from judgment_app.case_number import CaseNumber
from judgment_app.paths import resolve_output_path

IDENTITY_COLUMNS = {"裁判機關", "裁判年度", "裁判冠字", "裁判號數", "姓名"}


def load_comparison_report(path: str | Path) -> dict:
    """讀取既有報告，驗證顯示所需欄位並重新計算摘要；不修改檔案。"""
    source = resolve_output_path(path)
    report = json.loads(source.read_text(encoding="utf-8-sig"))
    try:
        return ComparisonReport.model_validate(report).to_dict()
    except ValidationError as exc:
        raise ValueError(f"不是有效的比對報告：{source}：{exc}") from exc


def compare_to_report(excel_path: Path, json_path: Path, destination: Path, *, compare, progress=None) -> dict:
    """執行指定的比對方法，回寫 Excel 結果並另存 JSON 報告。"""
    destination = resolve_output_path(destination).resolve()
    if destination.suffix.lower() != ".json":
        raise ValueError("比對報告必須為 .json 檔案")
    if any(Path(source).resolve() == destination for source in (excel_path, json_path)):
        raise ValueError("比對報告不可覆寫來源檔案")
    if destination.exists():
        data = json.loads(destination.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict) and "case_number" in data:
            raise ValueError("比對報告不可覆寫分析 JSON")
    report = compare(excel_path, json_path, progress=progress)
    if progress:
        progress(f"正在儲存比對報告：{destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def write_excel_statuses(path: Path, statuses: dict, *, legacy_tables=None) -> Path:
    """xlsx 原檔回寫；xls 以讀取的表格另存 xlsx，不覆寫既有結果檔。"""
    destination = path
    if legacy_tables is None:
        workbook = load_workbook(path)
    else:
        destination = path.with_name(f"{path.stem}.status.xlsx")
        number = 2
        while destination.exists():
            destination = path.with_name(f"{path.stem}.status-{number}.xlsx")
            number += 1
        workbook = Workbook()
        workbook.remove(workbook.active)
    temporary = None
    try:
        if legacy_tables is not None:
            for sheet, table in legacy_tables.items():
                worksheet = workbook.create_sheet(sheet)
                for row_number, row in enumerate(table.itertuples(index=False, name=None), 1):
                    for column, value in enumerate(row, 1):
                        if not pd.isna(value):
                            cell = worksheet.cell(row_number, column, value)
                            if isinstance(value, str):
                                cell.data_type = "s"
        for sheet, (header, row_statuses) in statuses.items():
            worksheet = workbook[sheet]
            for field in ("status", "reason"):
                column = next((cell.column for cell in worksheet[header]
                               if _text(cell.value) == field), worksheet.max_column + 1)
                worksheet.cell(header, column, field)
                for row_number in range(header + 1, worksheet.max_row + 1):
                    value = row_statuses.get(row_number, {}).get(field)
                    cell = worksheet.cell(row_number, column)
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = "s"
        with NamedTemporaryFile(dir=path.parent, prefix="~$status-", suffix=".xlsx", delete=False) as file:
            temporary = Path(file.name)
        workbook.save(temporary)
        temporary.replace(destination)
        return destination
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def matches_defendant_name(excel_name: str, json_name: str) -> bool:
    """遮蔽姓名依 ○ 前的任意非空前綴配對，保留遮蔽長度檢查；未遮蔽姓名完整比對。"""
    excel_name = _text(excel_name)
    json_name = _text(json_name)
    if "○" not in excel_name:
        return bool(excel_name) and excel_name == json_name
    excel_prefix, separator, excel_suffix = excel_name.partition("○")
    excel_suffix = separator + excel_suffix
    if not excel_prefix:
        return False
    json_suffix = json_name.split(excel_prefix, 1)[1] if excel_prefix in json_name else ""
    if len(excel_suffix) > len(json_suffix):
        return False
    return json_name.startswith(excel_prefix)


def find_excel_files(path: Path) -> list[Path]:
    files = find_files(path, {".xls", ".xlsx"})
    if path.is_dir():
        files = [file for file in files if not re.search(r"\.status(?:-\d+)?\.xlsx$", file.name, re.IGNORECASE)]
    if not files:
        raise ValueError(f"找不到 Excel：{path}")
    return files


def read_excel_header(table: pd.DataFrame) -> tuple[int, dict[str, int]]:
    """回傳標題列的零起算位置及欄位索引；同名取首欄，排除空白與舊結果欄。"""
    header = next((i for i, row in enumerate(table.itertuples(index=False, name=None))
                   if IDENTITY_COLUMNS <= {_text(value) for value in row}), None)
    if header is None:
        raise ValueError("找不到裁判案號及姓名標題列")
    columns = {}
    for i, name in enumerate(table.iloc[header]):
        columns.setdefault(_text(name), i)
    for name in ("", "status", "reason"):
        columns.pop(name, None)
    return header, columns


def iter_excel_rows(table: pd.DataFrame, header: int, columns: dict[str, int],
                    path: Path, sheet: str) -> Iterator[tuple[dict, dict]]:
    """逐列提供欄位值與實際 Excel 來源位置；全空白列不產生結果。"""
    for i in range(header + 1, len(table)):
        row = {name: table.iat[i, col] for name, col in columns.items()}
        if all(_is_blank(value) for value in row.values()):
            continue
        yield row, {"excel": str(path), "sheet": sheet, "row": i + 1}


def load_analyses(path: Path, errors: list) -> dict[CaseNumber, list[tuple[Path, JudgmentAnalysis]]]:
    index = {}
    for source in find_files(path, {".json"}):
        try:
            data = json.loads(source.read_text(encoding="utf-8-sig"))
            if not isinstance(data, dict) or "case_number" not in data:
                continue
            analysis = JudgmentAnalysis.from_dict(data)
            case_number = CaseNumber.from_string(analysis.case_number)
            index.setdefault(case_number, []).append((source, analysis))
        except (ValueError, TypeError, KeyError, OSError) as exc:
            errors.append({"file": str(source), "error": f"{type(exc).__name__}: {exc}"})
    return index


def find_files(path: Path, suffixes: set[str]) -> list[Path]:
    if not path.exists():
        raise ValueError(f"路徑不存在：{path}")
    return sorted(p for p in ([path] if path.is_file() else path.rglob("*"))
                  if p.is_file() and p.suffix.lower() in suffixes and not p.name.startswith("~$"))


def excel_bool(value, *, column: str, blank: bool | None = False) -> bool | None:
    """共用布林值解析；空白預設 False，可由選填欄位指定 None。"""
    text = _text(value)
    if not text:
        return blank
    bool_values = {
        "0": False, "0.0": False, "否": False, "n": False, "false": False,
        "1": True, "1.0": True, "是": True, "y": True, "true": True, "✓": True,
    }
    try:
        return bool_values[text.lower()]
    except KeyError:
        raise ValueError(f"{column}布林值無法辨識：{text}") from None


def excel_integer(value, *, zero_is_none: bool = False) -> int | None:
    if excel_optional_text(value) is None:
        return None
    number = _non_negative_integer(value)
    return None if zero_is_none and number == 0 else number


def excel_optional_text(value) -> str | None:
    """一般欄位的空白及 x 標記視為 None；布林欄位另依 excel_bool 解析。"""
    text = _text(value)
    return None if not text or text.lower() == "x" else text


def _non_negative_integer(value) -> int:
    number = float(value)
    if not number.is_integer() or number < 0:
        raise ValueError(f"非負整數格式錯誤：{value}")
    return int(number)


def _text(value) -> str:
    return "" if _is_blank(value) else re.sub(
        r"\s+", "", str(value)).replace("（", "(").replace("）", ")"
    )


def _is_blank(value) -> bool:
    return value is None or bool(pd.isna(value)) or str(value).strip() == ""
