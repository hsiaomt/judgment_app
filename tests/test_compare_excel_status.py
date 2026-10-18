from pathlib import Path
from tempfile import TemporaryDirectory
import json
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from judgment_app.analysis.compare import compare_files


class ExcelStatusTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.excel = self.root / 'source.xlsx'
        self.json_path = self.root / 'analysis.json'
        self.json_path.write_text(json.dumps({
            'case_number': '臺灣臺北地方法院114年度訴字第1號',
            'defendants': [{'name': '甲', 'is_recidivist': False,
                'crimes': [{'crime_id': 1, 'crime_jsst': '科刑', 'crime_name': '測試', 'times': 1,
                            'laws': [{'law_name': '刑法', 'article': 339,
                                      'is_attempted': False, 'is_accessory': False}]}],
                'execution_groups': [{'crime_ids': [1], 'executed_jsst': '科刑'}]}],
        }), encoding='utf-8')
        book = Workbook()
        sheet = book.active
        sheet.title = '比對'
        sheet.append(['原始標題'])
        sheet.merge_cells('A1:C1')
        sheet['A1'].font = Font(bold=True)
        sheet.append(['裁判機關', '裁判年度', '裁判冠字', '裁判號數', '姓名', '案由', '條', '累犯'])
        for name, year, recidivist in [('甲', 114, 0), ('甲', 114, 1), ('乙', 114, 0), ('甲', '錯誤', 0)]:
            sheet.append(['TPD', year, '訴', 1, name, '刑法', 339, recidivist])
        sheet.column_dimensions['A'].width = 25
        sheet.freeze_panes = 'A3'
        other = book.copy_worksheet(sheet)
        other.title = '第二張'
        other['E3'] = '乙'
        notes = book.create_sheet('說明')
        notes['A1'] = '=1+1'
        book.save(self.excel)
        book.close()

    def test_status_matches_report_and_preserves_workbook(self):
        report = compare_files(self.excel, self.json_path)
        book = load_workbook(self.excel)
        self.addCleanup(book.close)
        self.assertEqual([r['status'] for r in report['rows'][:4]],
                         ['equal', 'unmatched', 'unmatched', 'unmatched'])
        for row in report['rows']:
            self.assertEqual(book[row['sheet']].cell(row['row'], 9).value, row['status'])
        self.assertEqual(book['比對'].cell(2, 9).value, 'status')
        self.assertTrue(book['比對']['A1'].font.bold)
        self.assertIn('A1:C1', book['比對'].merged_cells)
        self.assertEqual(book['比對'].column_dimensions['A'].width, 25)
        self.assertEqual(book['比對'].freeze_panes, 'A3')
        self.assertEqual(book['說明']['A1'].value, '=1+1')
        self.assertEqual(book['說明'].max_column, 1)

    def test_rerun_updates_existing_status_and_clears_empty_row(self):
        compare_files(self.excel, self.json_path)
        book = load_workbook(self.excel)
        book['比對']['E3'] = '乙'
        for column in range(1, 9):
            book['比對'].cell(4, column).value = None
        book.save(self.excel)
        book.close()
        report = compare_files(self.excel, self.json_path)
        book = load_workbook(self.excel)
        self.addCleanup(book.close)
        self.assertEqual(book['比對'].max_column, 9)
        self.assertEqual(book['比對']['I3'].value, 'unmatched')
        self.assertIsNone(book['比對']['I4'].value)
        self.assertFalse(any(r['sheet'] == '比對' and r['row'] == 4 for r in report['rows']))
        self.assertNotIn('status', report['sheets'][0]['uncompared_columns'])

    def test_save_failure_preserves_source_and_reports_error(self):
        original = self.excel.read_bytes()
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=PermissionError('無法寫入')):
            report = compare_files(self.excel, self.json_path)
        self.assertEqual(self.excel.read_bytes(), original)
        self.assertTrue(any('無法寫入' in e['error'] for e in report['errors']))
        self.assertEqual(len(report['rows']), 8)
        self.assertEqual(list(self.root.glob('~$status-*')), [])

    def test_xls_exports_all_sheets_and_directory_skips_generated_files(self):
        source = self.root / 'legacy.xls'
        source.write_bytes(b'original legacy file')
        original_excel_file = pd.ExcelFile

        def read_book(path):
            # 用既有表格模擬 xlrd 讀取結果，驗證轉存與目錄掃描流程。
            return original_excel_file(self.excel if Path(path) == source else path)

        with patch('judgment_app.analysis.compare.pd.ExcelFile', side_effect=read_book):
            report = compare_files(source, self.json_path)
            output = Path(report['excel_outputs'][0]['output'])
            self.assertEqual(output.name, 'legacy.status.xlsx')
            first_output = output.read_bytes()
            book = load_workbook(output)
            self.assertEqual(book.sheetnames, ['比對', '第二張', '說明'])
            self.assertEqual(book['比對']['I2'].value, 'status')
            self.assertEqual(book['比對']['I3'].value, 'equal')
            book.close()
            report = compare_files(self.root, self.json_path)
            self.assertEqual(len(report['rows']), 16)  # 只讀原本的 xls 與 xlsx。
            self.assertEqual(len(report['excel_outputs']), 2)
            self.assertTrue((self.root / 'legacy.status-2.xlsx').is_file())
            self.assertEqual(output.read_bytes(), first_output)
            self.assertEqual(source.read_bytes(), b'original legacy file')
